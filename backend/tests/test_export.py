from datetime import date
from decimal import Decimal

import pytest
from app.models import Customer, Product, PurchaseDocument, PurchaseLine, SalesDocument, SalesLine
from app.services.export import build_workbook, export_workbook
from openpyxl import load_workbook

EXPECTED_SHEETS = [
    "Sales documents",
    "Sales lines",
    "Purchase documents",
    "Purchase lines",
    "Customers",
    "Products",
]


def _sales(session, source_id, document_date, total="1650000.00"):
    document = SalesDocument(
        source_system="easybooks",
        source_id=source_id,
        source_document_number=f"BH-{source_id}",
        document_date=document_date,
        accounting_object_code="KH-SYNTH-01",
        accounting_object_name="Synthetic Customer",
        currency_id="VND",
        subtotal=Decimal("1500000.00"),
        discount_amount=Decimal("0"),
        vat_amount=Decimal("150000.00"),
        total_amount=Decimal(total),
        normalized_hash=f"hash-{source_id}",
    )
    session.add(document)
    session.flush()
    session.add(
        SalesLine(
            sales_document_id=document.id,
            source_line_key=f"key-{source_id}",
            material_goods_code="PAPER-SYNTH-01",
            material_goods_name="Synthetic paper",
            unit="kg",
            quantity=Decimal("1000.0000"),
            unit_price=Decimal("1500.0000"),
            amount=Decimal("1500000.00"),
            discount_amount=Decimal("0"),
            vat_amount=Decimal("150000.00"),
        )
    )
    return document


def _purchase(session, source_id, document_date):
    document = PurchaseDocument(
        source_system="easybooks",
        source_id=source_id,
        source_document_number=f"MH-{source_id}",
        document_date=document_date,
        vendor_code="NCC-SYNTH-01",
        vendor_name="Synthetic Vendor",
        currency_rate=Decimal("1.000000"),
        total_purchase_amount=Decimal("2500000.00"),
        normalized_hash=f"phash-{source_id}",
    )
    session.add(document)
    session.flush()
    session.add(
        PurchaseLine(
            purchase_document_id=document.id,
            source_line_key=f"pkey-{source_id}",
            material_goods_code="SERVICE-SYNTH",
            material_goods_name="Synthetic service",
            unit="month",
            quantity=Decimal("1.0000"),
            unit_price=Decimal("0"),
            purchase_amount=Decimal("2500000.00"),
            discount_amount=Decimal("0"),
        )
    )
    return document


@pytest.fixture
def populated(session):
    _sales(session, "s-june", date(2026, 6, 1))
    _sales(session, "s-august", date(2026, 8, 1))
    _purchase(session, "p-june", date(2026, 6, 15))
    session.add(Customer(name="Synthetic Customer", easybooks_accounting_object_code="KH-SYNTH-01"))
    session.add(Product(code="PAPER-SYNTH-01", name="Synthetic paper", unit="kg"))
    session.commit()
    return session


def test_workbook_has_every_expected_sheet_in_order(populated):
    workbook, counts = build_workbook(populated)

    assert workbook.sheetnames == EXPECTED_SHEETS
    assert counts == {
        "Sales documents": 2,
        "Sales lines": 2,
        "Purchase documents": 1,
        "Purchase lines": 1,
        "Customers": 1,
        "Products": 1,
    }


def test_a_date_window_bounds_the_transactional_sheets(populated):
    _, counts = build_workbook(populated, from_date=date(2026, 7, 1), to_date=date(2026, 12, 31))

    assert counts["Sales documents"] == 1
    assert counts["Sales lines"] == 1
    # Catalog sheets are not date-scoped.
    assert counts["Customers"] == 1


def test_an_undated_document_is_never_silently_dropped_by_a_window(populated):
    """EasyBooks produces undated documents; excluding them would lose real rows."""
    _sales(populated, "s-undated", None)
    _purchase(populated, "p-undated", None)
    populated.commit()

    _, counts = build_workbook(populated, from_date=date(2026, 7, 1), to_date=date(2026, 7, 31))

    assert counts["Sales documents"] == 1
    assert counts["Purchase documents"] == 1
    assert counts["Sales lines"] == 1


def test_money_and_quantities_survive_as_exact_decimals(populated, tmp_path):
    export_workbook(populated, tmp_path / "out.xlsx")
    sheet = load_workbook(tmp_path / "out.xlsx")["Sales lines"]
    headers = [cell.value for cell in sheet[1]]

    amount = sheet.cell(row=2, column=headers.index("Amount") + 1)
    quantity = sheet.cell(row=2, column=headers.index("Quantity") + 1)

    assert Decimal(str(amount.value)) == Decimal("1500000.00")
    assert Decimal(str(quantity.value)) == Decimal("1000.0000")
    assert amount.number_format == "#,##0.00"
    assert quantity.number_format == "#,##0.0000"


def test_dates_are_written_as_dates_not_text(populated, tmp_path):
    export_workbook(populated, tmp_path / "out.xlsx")
    sheet = load_workbook(tmp_path / "out.xlsx")["Sales documents"]
    headers = [cell.value for cell in sheet[1]]

    cell = sheet.cell(row=2, column=headers.index("Document date") + 1)

    assert cell.value.date() == date(2026, 6, 1)
    assert cell.number_format == "yyyy-mm-dd"


def test_every_sheet_is_navigable(populated, tmp_path):
    export_workbook(populated, tmp_path / "out.xlsx")
    workbook = load_workbook(tmp_path / "out.xlsx")

    for name in EXPECTED_SHEETS:
        sheet = workbook[name]
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref is not None
        assert all(cell.font.bold for cell in sheet[1])


def test_an_empty_database_still_produces_labelled_sheets(session, tmp_path):
    summary = export_workbook(session, tmp_path / "empty.xlsx")
    workbook = load_workbook(tmp_path / "empty.xlsx")

    assert summary.total_rows == 0
    assert workbook.sheetnames == EXPECTED_SHEETS
    headers = [cell.value for cell in workbook["Customers"][1]]
    assert headers == ["EasyBooks code", "Name", "Tax code"]


def test_the_export_creates_missing_parent_directories(populated, tmp_path):
    summary = export_workbook(populated, tmp_path / "nested" / "deep" / "out.xlsx")

    assert summary.path.exists()
    assert summary.total_rows == 8


def test_the_export_reads_no_easybooks_api(populated, tmp_path, monkeypatch):
    """The export must work with no credential configured and no network call."""
    import httpx

    def explode(*args, **kwargs):
        raise AssertionError("the export must not perform any HTTP request")

    monkeypatch.setattr(httpx.Client, "request", explode)
    monkeypatch.setattr(httpx, "get", explode)
    monkeypatch.setattr(httpx, "post", explode)

    summary = export_workbook(populated, tmp_path / "out.xlsx")

    assert summary.path.exists()
