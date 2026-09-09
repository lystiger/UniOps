"""Export ingested UniOps data to a workbook.

This reads only from the UniOps database. It never contacts EasyBooks, so it
needs no credential and works on whatever the last sync stored.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from app.models import (
    Customer,
    Product,
    PurchaseDocument,
    PurchaseLine,
    SalesDocument,
    SalesLine,
)

# Excel renders a bare Decimal without separators, which makes money columns hard
# to scan. Quantities keep four places to match the NUMERIC(18,4) source columns.
MONEY_FORMAT = "#,##0.00"
QUANTITY_FORMAT = "#,##0.0000"
DATE_FORMAT = "yyyy-mm-dd"

MAX_COLUMN_WIDTH = 48


@dataclass
class SheetSpec:
    title: str
    headers: list[str]
    rows: list[list[Any]]


@dataclass
class ExportSummary:
    path: Path
    sheet_rows: dict[str, int]

    @property
    def total_rows(self) -> int:
        return sum(self.sheet_rows.values())


def _sales_documents(session: Session, from_date: date | None, to_date: date | None) -> SheetSpec:
    statement = select(SalesDocument).order_by(
        SalesDocument.document_date, SalesDocument.source_document_number
    )
    statement = _within(statement, SalesDocument.document_date, from_date, to_date)
    rows = [
        [
            document.source_id,
            document.source_document_number,
            document.invoice_series,
            document.invoice_number,
            document.document_date,
            document.posted_date,
            document.accounting_object_code,
            document.accounting_object_name,
            document.currency_id,
            document.subtotal,
            document.discount_amount,
            document.vat_amount,
            document.total_amount,
            document.recorded,
        ]
        for document in session.scalars(statement)
    ]
    return SheetSpec(
        "Sales documents",
        [
            "Source ID",
            "Document no.",
            "Invoice series",
            "Invoice no.",
            "Document date",
            "Posted date",
            "Customer code",
            "Customer name",
            "Currency",
            "Subtotal",
            "Discount",
            "VAT",
            "Total",
            "Recorded",
        ],
        rows,
    )


def _sales_lines(session: Session, from_date: date | None, to_date: date | None) -> SheetSpec:
    statement = (
        select(SalesLine, SalesDocument)
        .join(SalesDocument, SalesLine.sales_document_id == SalesDocument.id)
        .order_by(SalesDocument.document_date, SalesDocument.source_document_number)
    )
    statement = _within(statement, SalesDocument.document_date, from_date, to_date)
    rows = [
        [
            document.source_id,
            document.source_document_number,
            document.document_date,
            line.material_goods_code,
            line.material_goods_name,
            line.unit,
            line.quantity,
            line.unit_price,
            line.amount,
            line.discount_amount,
            line.vat_amount,
            line.repository_code,
        ]
        for line, document in session.execute(statement)
    ]
    return SheetSpec(
        "Sales lines",
        [
            "Document source ID",
            "Document no.",
            "Document date",
            "Product code",
            "Product name",
            "Unit",
            "Quantity",
            "Unit price",
            "Amount",
            "Discount",
            "VAT",
            "Warehouse",
        ],
        rows,
    )


def _purchase_documents(
    session: Session, from_date: date | None, to_date: date | None
) -> SheetSpec:
    statement = select(PurchaseDocument).order_by(
        PurchaseDocument.document_date, PurchaseDocument.source_document_number
    )
    statement = _within(statement, PurchaseDocument.document_date, from_date, to_date)
    rows = [
        [
            document.source_id,
            document.source_document_number,
            document.invoice_number,
            document.document_date,
            document.posted_date,
            document.vendor_code,
            document.vendor_name,
            document.tax_code,
            document.currency_rate,
            document.total_purchase_amount,
        ]
        for document in session.scalars(statement)
    ]
    return SheetSpec(
        "Purchase documents",
        [
            "Source ID",
            "Document no.",
            "Invoice no.",
            "Document date",
            "Posted date",
            "Vendor code",
            "Vendor name",
            "Tax code",
            "Currency rate",
            "Total purchase",
        ],
        rows,
    )


def _purchase_lines(session: Session, from_date: date | None, to_date: date | None) -> SheetSpec:
    statement = (
        select(PurchaseLine, PurchaseDocument)
        .join(PurchaseDocument, PurchaseLine.purchase_document_id == PurchaseDocument.id)
        .order_by(PurchaseDocument.document_date, PurchaseDocument.source_document_number)
    )
    statement = _within(statement, PurchaseDocument.document_date, from_date, to_date)
    rows = [
        [
            document.source_id,
            document.source_document_number,
            document.document_date,
            line.material_goods_code,
            line.material_goods_name,
            line.unit,
            line.quantity,
            line.unit_price,
            line.purchase_amount,
            line.discount_amount,
            line.vat_amount,
            line.warehouse_code,
            line.description,
        ]
        for line, document in session.execute(statement)
    ]
    return SheetSpec(
        "Purchase lines",
        [
            "Document source ID",
            "Document no.",
            "Document date",
            "Item code",
            "Item name",
            "Unit",
            "Quantity",
            "Unit price",
            "Purchase amount",
            "Discount",
            "VAT amount",
            "Warehouse",
            "Description",
        ],
        rows,
    )


def _customers(session: Session) -> SheetSpec:
    rows = [
        [customer.easybooks_accounting_object_code, customer.name, customer.tax_code]
        for customer in session.scalars(select(Customer).order_by(Customer.name))
    ]
    return SheetSpec("Customers", ["EasyBooks code", "Name", "Tax code"], rows)


def _products(session: Session) -> SheetSpec:
    rows = [
        [product.code, product.name, product.unit, product.easybooks_material_goods_id]
        for product in session.scalars(select(Product).order_by(Product.code))
    ]
    return SheetSpec("Products", ["Code", "Name", "Unit", "EasyBooks material ID"], rows)


def _within(statement: Any, column: Any, from_date: date | None, to_date: date | None) -> Any:
    """Bound a statement by document date, keeping rows that carry no date.

    An undated document cannot be shown to fall outside the window, and EasyBooks
    does produce them: one live purchase document has no date at all. Excluding it
    would drop a real record from the export with nothing to indicate the loss, so
    undated rows are always kept.
    """
    if from_date is not None:
        statement = statement.where(or_(column.is_(None), column >= from_date))
    if to_date is not None:
        statement = statement.where(or_(column.is_(None), column <= to_date))
    return statement


def _write_sheet(workbook: Workbook, spec: SheetSpec, *, first: bool) -> None:
    sheet = workbook.active if first else workbook.create_sheet()
    sheet.title = spec.title
    sheet.append(spec.headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")
    for row in spec.rows:
        sheet.append(row)

    for index, header in enumerate(spec.headers, start=1):
        letter = get_column_letter(index)
        widest = max(
            [len(str(header))] + [len(str(row[index - 1] or "")) for row in spec.rows] or [0]
        )
        sheet.column_dimensions[letter].width = min(widest + 2, MAX_COLUMN_WIDTH)
        for cell in sheet[letter][1:]:
            if isinstance(cell.value, Decimal):
                cell.number_format = QUANTITY_FORMAT if _is_quantity(header) else MONEY_FORMAT
            elif isinstance(cell.value, date) and not isinstance(cell.value, datetime):
                cell.number_format = DATE_FORMAT

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _is_quantity(header: str) -> bool:
    return header in {"Quantity", "Currency rate"}


def build_workbook(
    session: Session, *, from_date: date | None = None, to_date: date | None = None
) -> tuple[Workbook, dict[str, int]]:
    specs = [
        _sales_documents(session, from_date, to_date),
        _sales_lines(session, from_date, to_date),
        _purchase_documents(session, from_date, to_date),
        _purchase_lines(session, from_date, to_date),
        _customers(session),
        _products(session),
    ]
    workbook = Workbook()
    for position, spec in enumerate(specs):
        _write_sheet(workbook, spec, first=position == 0)
    return workbook, {spec.title: len(spec.rows) for spec in specs}


def export_workbook(
    session: Session,
    destination: Path,
    *,
    from_date: date | None = None,
    to_date: date | None = None,
) -> ExportSummary:
    workbook, sheet_rows = build_workbook(session, from_date=from_date, to_date=to_date)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return ExportSummary(destination, sheet_rows)
